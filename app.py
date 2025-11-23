"""
EthicCheck - AI-Powered Ethical Analysis for Student Projects
Complete Streamlit Application with Groq API Integration + GitHub Analysis + Reset Button
"""

import streamlit as st
import os
import re
import json
from groq import Groq
from sentence_transformers import SentenceTransformer
import numpy as np
from datetime import datetime
import PyPDF2
import io
from dotenv import load_dotenv
import pandas as pd
from fpdf import FPDF
import docx

# ==================== CUSTOM MODULE IMPORTS ====================
from bias_fairness_checker import (
    run_bias_check, 
    generate_bias_summary,
    parse_uploaded_dataset
)

from ip_copyright_checker import (
    analyze_ip_copyright,
    generate_ip_summary
)

from privacy_checker import (
    run_privacy_check,
    generate_privacy_summary
)

from plagiarism_checker import (
    run_plagiarism_check,
    generate_plagiarism_summary
)

from github_analyzer import (
    analyze_github_repository,
    is_valid_github_url,
    format_repo_info_display
)

# ==================== CONFIGURATION ====================

# Page configuration
st.set_page_config(
    page_title="EthicCheck - AI Ethics Analyzer",
    page_icon="🛡️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom CSS for professional styling
st.markdown("""
<style>
    .main-header {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        padding: 2rem;
        border-radius: 10px;
        color: white;
        margin-bottom: 2rem;
    }
    .metric-card {
        background: white;
        padding: 1.5rem;
        border-radius: 8px;
        border-left: 4px solid #667eea;
        box-shadow: 0 2px 4px rgba(0,0,0,0.1);
    }
    .issue-card-high {
        background: #fee2e2;
        border-left: 4px solid #dc2626;
        padding: 1rem;
        border-radius: 8px;
        margin: 1rem 0;
    }
    .issue-card-medium {
        background: #fef3c7;
        border-left: 4px solid #f59e0b;
        padding: 1rem;
        border-radius: 8px;
        margin: 1rem 0;
    }
    .issue-card-low {
        background: #dbeafe;
        border-left: 4px solid #3b82f6;
        padding: 1rem;
        border-radius: 8px;
        margin: 1rem 0;
    }
    .stButton>button {
        width: 100%;
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        color: white;
        border: none;
        padding: 0.75rem 2rem;
        font-weight: 600;
        border-radius: 8px;
    }
    div[data-testid="column"] > div > div.reset-button button {
        background: linear-gradient(135deg, #ef4444 0%, #dc2626 100%) !important;
        color: white !important;
        border: none !important;
        padding: 0.6rem 1.5rem !important;
        font-weight: 600 !important;
        border-radius: 8px !important;
        transition: all 0.3s ease !important;
    }
    div[data-testid="column"] > div > div.reset-button button:hover {
        background: linear-gradient(135deg, #dc2626 0%, #b91c1c 100%) !important;
        transform: translateY(-2px);
        box-shadow: 0 4px 12px rgba(239, 68, 68, 0.4) !important;
    }
    .repo-info-box {
        background: #f8f9fa;
        border-left: 4px solid #667eea;
        padding: 1rem;
        border-radius: 8px;
        margin: 1rem 0;
    }
</style>
""", unsafe_allow_html=True)

# ==================== INITIALIZATION ====================

# Initialize session state
if 'analysis_results' not in st.session_state:
    st.session_state.analysis_results = None
if 'active_tab' not in st.session_state:
    st.session_state.active_tab = 0
if 'repo_info' not in st.session_state:
    st.session_state.repo_info = None
if 'reset_counter' not in st.session_state:
    st.session_state.reset_counter = 0

# Initialize Groq client
@st.cache_resource
def init_groq_client():
    load_dotenv()
    try:
        api_key = st.secrets.get("GROQ_API_KEY", "")
    except:
        api_key = os.getenv("GROQ_API_KEY", "")
    if not api_key:
        st.warning("⚠️ Please set your GROQ_API_KEY environment variable")
        return None
    return Groq(api_key=api_key)

# Initialize embedding model
@st.cache_resource
def init_embedding_model():
    return SentenceTransformer('all-MiniLM-L6-v2')

groq_client = init_groq_client()
embedding_model = init_embedding_model()


# ==================== UTILITY FUNCTIONS ====================

def extract_text_from_pdf(pdf_file):
    """Extract text from uploaded PDF"""
    try:
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(pdf_file.read()))
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text
    except Exception as e:
        return f"Error extracting PDF: {str(e)}"

def extract_text_from_docx(docx_file):
    """Extract text from uploaded DOCX"""
    try:
        doc = docx.Document(io.BytesIO(docx_file.read()))
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text
    except Exception as e:
        return f"Error extracting DOCX: {str(e)}"

def detect_code_issues(code_text):
    """Detect security issues in code"""
    issues = []
    
    # Check for dangerous operations
    dangerous_ops = ['os.system', 'subprocess.call', 'eval(', 'exec(']
    for op in dangerous_ops:
        if op in code_text:
            issues.append({
                'type': f'Dangerous Operation: {op}',
                'severity': 'medium',
                'count': code_text.count(op)
            })
    
    return issues

def check_ethical_keywords(text):
    """Check for ethical red flags"""
    red_flags = []
    
    high_risk_keywords = [
        'scrape', 'bypass', 'circumvent', 'crack', 'exploit',
        'phishing', 'surveillance', 'deepfake', 'weaponize'
    ]
    
    medium_risk_keywords = [
        'face recognition', 'biometric', 'tracking', 'profile',
        'without consent', 'unauthorized', 'private data'
    ]
    
    text_lower = text.lower()
    
    for keyword in high_risk_keywords:
        if keyword in text_lower:
            context = extract_context(text, keyword)
            red_flags.append({
                'keyword': keyword,
                'severity': 'high',
                'context': context
            })
    
    for keyword in medium_risk_keywords:
        if keyword in text_lower:
            context = extract_context(text, keyword)
            red_flags.append({
                'keyword': keyword,
                'severity': 'medium',
                'context': context
            })
    
    return red_flags

def extract_context(text, keyword, window=100):
    """Extract context around keyword"""
    idx = text.lower().find(keyword.lower())
    if idx == -1:
        return ""
    start = max(0, idx - window)
    end = min(len(text), idx + len(keyword) + window)
    return "..." + text[start:end] + "..."

def detect_artifact_type(text, uploaded_df=None):
    """Auto-detect project type from content"""
    if uploaded_df is not None:
        return "Dataset Description"
    
    text_lower = text.lower()
    
    # Check for code indicators
    code_indicators = ['import ', 'def ', 'class ', 'function ', '#!/usr/bin', 'const ', 'var ', 'let ']
    if any(indicator in text_lower for indicator in code_indicators):
        return "Code"
    
    # Check for dataset indicators
    dataset_indicators = ['dataset', 'data collection', 'sample size', 'participants', 'variables']
    if sum(indicator in text_lower for indicator in dataset_indicators) >= 2:
        return "Dataset Description"
    
    # Check for proposal indicators
    proposal_indicators = ['abstract', 'introduction', 'methodology', 'objectives', 'research question']
    if sum(indicator in text_lower for indicator in proposal_indicators) >= 2:
        return "Proposal"
    
    return "Full Report"

# ==================== PDF EXPORT FUNCTION ====================

def clean_text_for_pdf(text):
    """Clean text to remove characters not supported by latin-1 encoding."""
    if not text:
        return ""
    
    # Replace common unicode characters with ASCII equivalents
    replacements = {
        '\u2013': '-',  # en dash
        '\u2014': '--',  # em dash
        '\u2018': "'",  # left single quote
        '\u2019': "'",  # right single quote
        '\u201c': '"',  # left double quote
        '\u201d': '"',  # right double quote
        '\u2022': '*',  # bullet
        '\u2026': '...',  # ellipsis
        '\u00a0': ' ',  # non-breaking space
        '\u2192': '->',  # right arrow
        '\u00b7': '*',  # middle dot
        '\u25aa': '*',  # small square
        '\u2022': '*',  # bullet point
    }
    
    for unicode_char, ascii_char in replacements.items():
        text = text.replace(unicode_char, ascii_char)
    
    # Remove any remaining non-latin-1 characters
    text = text.encode('latin-1', errors='ignore').decode('latin-1')
    
    return text

def generate_pdf_report(results, bias_findings, ip_findings, privacy_findings, plagiarism_findings):
    """Generate PDF report of analysis results."""
    
    class PDF(FPDF):
        def header(self):
            self.set_font('Arial', 'B', 16)
            self.set_fill_color(102, 126, 234)
            self.set_text_color(255, 255, 255)
            self.cell(0, 15, 'EthicCheck Analysis Report', 0, 1, 'C', fill=True)
            self.ln(5)
        
        def footer(self):
            self.set_y(-15)
            self.set_font('Arial', 'I', 8)
            self.set_text_color(128)
            self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')
        
        def chapter_title(self, title):
            self.set_font('Arial', 'B', 14)
            self.set_fill_color(240, 240, 240)
            title = clean_text_for_pdf(title)
            self.cell(0, 10, title, 0, 1, 'L', fill=True)
            self.ln(2)
        
        def chapter_body(self, body):
            self.set_font('Arial', '', 11)
            body = clean_text_for_pdf(body)
            self.multi_cell(0, 6, body)
            self.ln()
    
    pdf = PDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    # Report metadata
    pdf.set_font('Arial', '', 10)
    pdf.set_text_color(100)
    pdf.cell(0, 6, f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}', 0, 1)
    pdf.cell(0, 6, 'Powered by EthicCheck - AI Ethics Analyzer', 0, 1)
    pdf.ln(5)
    
    # Overall Summary
    pdf.chapter_title('Overall Analysis Summary')
    issues = results.get('issues', [])
    high_count = len([i for i in issues if i.get('severity') == 'high'])
    medium_count = len([i for i in issues if i.get('severity') == 'medium'])
    low_count = len([i for i in issues if i.get('severity') == 'low'])
    
    summary_text = f"Overall Risk Level: {results.get('overall_score', 'N/A').upper()}\n"
    summary_text += f"Total Issues Found: {len(issues)}\n"
    summary_text += f"  - High Priority: {high_count}\n"
    summary_text += f"  - Medium Priority: {medium_count}\n"
    summary_text += f"  - Low Priority: {low_count}\n"
    
    if st.session_state.get('plagiarism_percentage') is not None:
        plag_pct = st.session_state.get('plagiarism_percentage', 0)
        summary_text += f"\nPlagiarism Score: {plag_pct}%\n"
    
    pdf.chapter_body(summary_text)
    
    # Student Summary
    pdf.chapter_title('Summary for Students')
    student_summary = results.get('student_summary', 'No summary available')
    pdf.chapter_body(student_summary)
    
    # Detailed Issues
    pdf.chapter_title('Detailed Issues')
    
    if not issues:
        pdf.chapter_body('No significant ethical issues detected!')
    else:
        # Sort by severity
        sorted_issues = sorted(
            issues, 
            key=lambda x: {'high': 0, 'medium': 1, 'low': 2}.get(x.get('severity', 'low'), 2)
        )
        
        for idx, issue in enumerate(sorted_issues, 1):
            severity = issue.get('severity', 'low')
            severity_icon = {'high': '[HIGH]', 'medium': '[MEDIUM]', 'low': '[LOW]'}.get(severity, '[INFO]')
            
            pdf.set_font('Arial', 'B', 12)
            pdf.set_text_color(0)
            issue_title = clean_text_for_pdf(issue.get('title', 'Issue'))
            pdf.cell(0, 8, f"{idx}. {severity_icon} {issue_title}", 0, 1)
            
            pdf.set_font('Arial', '', 10)
            pdf.set_text_color(50)
            
            # Category and Location
            category = clean_text_for_pdf(issue.get('category', 'N/A'))
            location = clean_text_for_pdf(issue.get('location', 'N/A'))
            pdf.cell(0, 6, f"Category: {category}", 0, 1)
            pdf.cell(0, 6, f"Location: {location}", 0, 1)
            
            # Evidence
            pdf.set_font('Arial', 'B', 10)
            pdf.cell(0, 6, 'Evidence:', 0, 1)
            pdf.set_font('Arial', '', 9)
            evidence = clean_text_for_pdf(issue.get('evidence', 'N/A'))
            if len(evidence) > 300:
                evidence = evidence[:300] + '...'
            pdf.multi_cell(0, 5, evidence)
            
            # Recommendation
            pdf.set_font('Arial', 'B', 10)
            pdf.cell(0, 6, 'Recommendation:', 0, 1)
            pdf.set_font('Arial', '', 9)
            recommendation = clean_text_for_pdf(issue.get('recommendation', 'N/A'))
            pdf.multi_cell(0, 5, recommendation)
            
            # Suggested Fix
            suggested_fix = issue.get('suggested_rewrite')
            if not suggested_fix and issue.get('category') == 'IP & Copyright':
                suggested_fix = generate_ip_fix_with_groq(issue)
            
            if suggested_fix:
                pdf.set_font('Arial', 'B', 10)
                pdf.cell(0, 6, 'Suggested Fix:', 0, 1)
                pdf.set_font('Arial', '', 9)
                suggested_fix = clean_text_for_pdf(suggested_fix)
                if len(suggested_fix) > 300:
                    suggested_fix = suggested_fix[:300] + '...'
                pdf.multi_cell(0, 5, suggested_fix)
            
            pdf.ln(3)
            pdf.set_draw_color(200)
            pdf.line(10, pdf.get_y(), 200, pdf.get_y())
            pdf.ln(3)
    
    # Instructor Notes
    pdf.add_page()
    pdf.chapter_title('Instructor Notes')
    instructor_notes = clean_text_for_pdf(results.get('instructor_notes', 'No additional notes'))
    pdf.chapter_body(instructor_notes)
    
    # Generate PDF in memory
    pdf_output = pdf.output(dest='S').encode('latin-1')
    return pdf_output

# ==================== GROQ FIX GENERATION ====================

def generate_ip_fix_with_groq(issue: dict) -> str:
    """Generate AI-powered fix suggestions for IP & Copyright issues using Groq."""
    if not groq_client:
        return "Enable Groq API for AI-powered fix suggestions."
    
    try:
        prompt = f"""You are an IP & Copyright compliance expert. Suggest a specific fix for this issue:

Issue: {issue.get('title', 'Unknown Issue')}
Evidence: {issue.get('evidence', 'No evidence provided')}
Current Recommendation: {issue.get('recommendation', 'No recommendation')}

Provide a concrete, actionable fix in 2-3 sentences. Be specific about what the student should do."""

        response = groq_client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[
                {"role": "system", "content": "You are a helpful IP & Copyright compliance assistant. Provide clear, actionable advice."},
                {"role": "user", "content": prompt}
            ],
            temperature=0.3,
            max_tokens=200
        )
        
        return response.choices[0].message.content.strip()
    except Exception as e:
        return f"Error generating fix: {str(e)}"

# ==================== GROQ API INTEGRATION ====================

def analyze_with_groq(text, artifact_type, check_options, uploaded_df=None):
    """Main analysis function using Groq API"""
    
    if not groq_client:
        return None
    
    # Build analysis prompt
    prompt = build_analysis_prompt(text, artifact_type, check_options)
    
    try:
        response = groq_client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[
                {
                    "role": "system",
                    "content": """You are EthicCheck, an AI assistant specialized in analyzing student projects for ethical concerns. 
                    You identify issues related to privacy, bias, licensing, plagiarism, security, and harmful use.
                    Always output valid JSON with the following structure:
                    {
                        "overall_score": "low|medium|high",
                        "issues": [
                            {
                                "category": "Privacy|Bias|License|Plagiarism|Security|Harmful Use",
                                "severity": "low|medium|high",
                                "title": "Brief title",
                                "evidence": "Quote from text",
                                "location": "Section reference",
                                "recommendation": "Concrete fix",
                                "suggested_rewrite": "Optional rewritten text"
                            }
                        ],
                        "student_summary": "Brief summary for student",
                        "instructor_notes": "Notes for instructor review"
                    }"""
                },
                {
                    "role": "user",
                    "content": prompt
                }
            ],
            temperature=0.3,
            max_tokens=2000
        )
        
        result = response.choices[0].message.content
        
        # Try to parse JSON
        try:
            if "```json" in result:
                result = result.split("```json")[1].split("```")[0]
            elif "```" in result:
                result = result.split("```")[1].split("```")[0]
            
            return json.loads(result)
        except json.JSONDecodeError:
            return {
                "overall_score": "medium",
                "issues": [{
                    "category": "Analysis",
                    "severity": "medium",
                    "title": "Analysis Completed",
                    "evidence": "See detailed response",
                    "location": "Full document",
                    "recommendation": result,
                    "suggested_rewrite": None
                }],
                "student_summary": result[:500],
                "instructor_notes": "Please review the full analysis."
            }
            
    except Exception as e:
        st.error(f"Error calling Groq API: {str(e)}")
        return None

def build_analysis_prompt(text, artifact_type, check_options):
    """Build structured prompt for Groq"""
    
    prompt = f"""Analyze the following {artifact_type} for ethical concerns.

ARTIFACT TEXT:
{text[:4000]}

ANALYSIS FOCUS:
"""
    
    if check_options.get('copyright', False):
        prompt += "- Intellectual Property and Copyright issues\n"
    if check_options.get('privacy', False):
        prompt += "- Privacy and PII exposure\n"
    if check_options.get('bias', False):
        prompt += "- Bias and fairness issues\n"
    if check_options.get('plagiarism', False):
        prompt += "- Potential plagiarism and attribution\n"
    
    prompt += """
POLICY RULES:
- Copyright: Using copyrighted material without permission = HIGH severity
- Privacy: Personal data without consent = HIGH severity
- Bias: Underrepresented groups in datasets = MEDIUM severity
- Plagiarism: Unattributed work or high similarity = HIGH severity

Output valid JSON only. Be specific with evidence and provide actionable recommendations.
"""
    
    return prompt

# ==================== UI COMPONENTS ====================

def render_header():
    """Render main header"""
    st.markdown("""
    <div class="main-header">
        <h1>🛡️ EthicCheck</h1>
        <p style="font-size: 1.2rem; margin-top: 0.5rem;">
            AI-Powered Ethical Analysis for Student Projects
        </p>
        <p style="font-size: 0.9rem; opacity: 0.9; margin-top: 0.5rem;">
            Powered by Groq • Llama 3 • Sentence Transformers
        </p>
    </div>
    """, unsafe_allow_html=True)

def render_upload_section():
    """Render upload and input section"""
    st.subheader("📤 Submit Your Project")
    
    col1, col2 = st.columns([2, 1])
    
    with col1:
        # File uploader with unique key based on reset counter
        uploaded_file = st.file_uploader(
            "Upload document or code",
            type=['txt', 'pdf', 'py', 'csv', 'docx'],
            help="Supported: PDF, TXT, Python, CSV, DOCX",
            key=f"file_uploader_{st.session_state.reset_counter}"
        )

        if uploaded_file is not None:
            st.success(f"✅ File uploaded: {uploaded_file.name} ({uploaded_file.size} bytes)")
        
        # Text area with unique key based on reset counter
        text_input = st.text_area(
            "Or paste your content here",
            height=200,
            placeholder="Paste your project proposal, code, or methodology...",
            key=f"text_input_{st.session_state.reset_counter}"
        )
        
        # GitHub URL with unique key based on reset counter
        git_url = st.text_input(
            "Or enter GitHub repository URL",
            placeholder="https://github.com/username/repo",
            help="Enter a public GitHub repository URL. Max size: 50MB, 500 files",
            key=f"git_url_{st.session_state.reset_counter}"
        )
        
        # Show URL validation
        if git_url and not is_valid_github_url(git_url):
            st.warning("⚠️ Invalid GitHub URL format. Use: https://github.com/user/repo")
    
    with col2:
        st.markdown("### Analysis Options")
        check_options = {
            'copyright': st.checkbox("IP & Copyright", value=True),
            'privacy': st.checkbox("Privacy Check", value=True),
            'bias': st.checkbox("Bias & Fairness", value=True),
            'plagiarism': st.checkbox("Plagiarism Check", value=True)
        }
        
        # Warnings
        if check_options['plagiarism']:
            st.caption("⏱️ Plagiarism check may take 20-40 seconds")
        
        if git_url:
            st.caption("📦 Repository analysis includes all code files")
        
        # Reset button with custom styling
        # In render_upload_section(), find the reset button section and replace with:

        # Reset button with custom styling
        st.markdown("---")
        st.markdown('<div class="reset-button">', unsafe_allow_html=True)
        if st.button(
            "🔄 Reset All Inputs", 
            use_container_width=True, 
            help="Clear all inputs and results",
            key=f"reset_btn_{st.session_state.reset_counter}"
        ):
            st.session_state.reset_counter += 1
            st.session_state.analysis_results = None
            st.session_state.bias_findings = []
            st.session_state.ip_findings = []
            st.session_state.privacy_findings = []
            st.session_state.plagiarism_findings = []
            st.session_state.plagiarism_percentage = None
            st.session_state.plagiarism_sources = []
            st.session_state.repo_info = None
            st.rerun()
        st.markdown('</div>', unsafe_allow_html=True)
    
    return uploaded_file, text_input, git_url, check_options

def render_results(results, bias_findings=None, ip_findings=None, privacy_findings=None, plagiarism_findings=None):
    """Render analysis results"""
    if not results:
        return
    
    # Overall metrics
    st.markdown("### 📊 Analysis Summary")
    
    col1, col2, col3, col4 = st.columns(4)
    
    issues = results.get('issues', [])
    high_count = len([i for i in issues if i.get('severity') == 'high'])
    medium_count = len([i for i in issues if i.get('severity') == 'medium'])
    low_count = len([i for i in issues if i.get('severity') == 'low'])
    
    with col1:
        st.metric("Overall Risk", results.get('overall_score', 'N/A').upper())
    with col2:
        st.metric("High Priority", high_count, delta=None, delta_color="inverse")
    with col3:
        st.metric("Medium Priority", medium_count)
    with col4:
        st.metric("Low Priority", low_count)
    
    # Plagiarism score if available
    if st.session_state.get('plagiarism_percentage') is not None:
        plag_pct = st.session_state.get('plagiarism_percentage', 0)
        if plag_pct >= 60:
            plag_emoji = "🔴"
            plag_level = "CRITICAL"
        elif plag_pct >= 40:
            plag_emoji = "🟠"
            plag_level = "HIGH"
        elif plag_pct >= 20:
            plag_emoji = "🟡"
            plag_level = "MEDIUM"
        elif plag_pct >= 10:
            plag_emoji = "🟢"
            plag_level = "LOW"
        else:
            plag_emoji = "✅"
            plag_level = "MINIMAL"
        st.markdown(f"**{plag_emoji} Plagiarism Score:** {plag_pct}% ({plag_level})")
    
    # Student summary
    st.markdown("### 📝 Summary for Students")
    st.info(results.get('student_summary', 'No summary available'))
    
    # Issues details
    st.markdown("### 🔍 Detailed Issues")
    
    if not issues:
        st.success("✅ No significant ethical issues detected!")
        return
    
    # Sort by severity
    sorted_issues = sorted(
        issues, 
        key=lambda x: {'high': 0, 'medium': 1, 'low': 2}.get(x.get('severity', 'low'), 2)
    )
    
    for idx, issue in enumerate(sorted_issues):
        severity = issue.get('severity', 'low')
        
        with st.expander(
            f"{'🔴' if severity == 'high' else '🟡' if severity == 'medium' else '🔵'} "
            f"{issue.get('title', 'Issue')} [{severity.upper()}]"
        ):
            st.markdown(f"**Category:** {issue.get('category', 'N/A')}")
            st.markdown(f"**Location:** {issue.get('location', 'N/A')}")
            
            st.markdown("**Evidence:**")
            st.code(issue.get('evidence', 'N/A'), language=None)
            
            st.markdown("**Recommendation:**")
            st.write(issue.get('recommendation', 'N/A'))
            
            # Handle suggested_rewrite - generate if missing for IP & Copyright
            suggested_fix = issue.get('suggested_rewrite')
            if not suggested_fix and issue.get('category') == 'IP & Copyright':
                # Generate fix using Groq
                suggested_fix = generate_ip_fix_with_groq(issue)
            
            if suggested_fix:
                st.markdown("**Suggested Fix:**")
                st.success(suggested_fix)
                
                unique_key = f"copy_{idx}_{issue.get('category', 'unknown')}_{severity}"
                if st.button(f"Copy Fix #{idx+1}", key=unique_key):
                    st.write("✅ Copied to clipboard (simulated)")
    
    # Instructor notes
    with st.expander("👨‍🏫 Instructor Notes"):
        st.write(results.get('instructor_notes', 'No additional notes'))
    
    # Export options
    st.markdown("### 💾 Export")
    
    # Generate PDF button
    if st.button("📄 Export PDF Report", use_container_width=True):
        with st.spinner("Generating PDF report..."):
            try:
                pdf_bytes = generate_pdf_report(
                    results, 
                    bias_findings or [], 
                    ip_findings or [], 
                    privacy_findings or [], 
                    plagiarism_findings or []
                )
                
                # Create download button
                st.download_button(
                    label="📥 Download PDF Report",
                    data=pdf_bytes,
                    file_name=f"EthicCheck_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )
                st.success("✅ PDF report generated successfully!")
            except Exception as e:
                st.error(f"Error generating PDF: {str(e)}")
                st.info("Note: Make sure fpdf is installed: pip install fpdf")

# ==================== MAIN APP ====================

def main():
    render_header()
    
    # Sidebar
    with st.sidebar:
        st.markdown("### ℹ️ About EthicCheck")
        st.write("""
        **AI-Powered Ethical Analysis Tool**
        
        EthicCheck helps students and educators identify potential ethical issues in academic projects.
        
        **Analysis Categories:**
        - 📜 IP & Copyright Compliance
        - 🔒 Privacy & Data Protection
        - ⚖️ Bias & Fairness Assessment
        - 🔍 Plagiarism Detection
        """)
        
        st.markdown("---")
        st.markdown("### 🎯 How It Works")
        st.write("""
        1. Upload your project, paste text, or enter GitHub URL
        2. Select analysis options
        3. Get instant AI-powered feedback
        4. Review issues and apply fixes
        """)
        
        st.markdown("---")
        st.markdown("### 📥 Input Priority")
        st.write("""
        **Priority Order:**
        1. 🔗 GitHub URL (highest)
        2. 📄 File Upload
        3. 📝 Paste Text (lowest)
        
        **Supported Files:**
        - PDF documents
        - TXT files
        - Python code (.py)
        - CSV datasets
        - DOCX documents
        """)
        
        st.markdown("---")
        st.markdown("### 🔐 Privacy & Security")
        st.write("""
        - ✅ Ephemeral processing
        - ✅ No permanent storage
        - ✅ FERPA compliant
        - ✅ Powered by Groq AI
        """)
        
        st.markdown("---")
        st.markdown("### 💡 Tips")
        st.write("""
        - GitHub repos: Max 50MB, 500 files
        - Review all high-priority issues
        - Apply suggested fixes
        - Re-analyze after changes
        - Use Reset button to start over
        """)
    
    # Main content
    tab1, tab2 = st.tabs(["📤 Upload & Analyze", "📊 Results"])
    
    with tab1:
        uploaded_file, text_input, git_url, check_options = render_upload_section()
        
        if st.button("🚀 Analyze Project", type="primary"):
            # Get input text and dataset
            input_text = ""
            uploaded_df = None
            is_github_repo = False
            
            # Priority: GitHub URL > File Upload > Text Input
            if git_url and is_valid_github_url(git_url):
                is_github_repo = True
                
                with st.spinner("🔄 Cloning and analyzing repository..."):
                    progress_placeholder = st.empty()
                    
                    def github_progress(message):
                        progress_placeholder.info(f"📦 {message}")
                    
                    # Analyze GitHub repository
                    repo_result = analyze_github_repository(git_url, github_progress)
                    
                    if repo_result['status'] == 'error':
                        st.error(f"❌ {repo_result['error']}")
                        return
                    
                    # Store repository info
                    st.session_state.repo_info = repo_result['repo_info']
                    
                    # Get combined content for analysis
                    content = repo_result['content']
                    input_text = content['code'] + "\n\n" + content['docs']
                    
                    # Success message
                    progress_placeholder.empty()
                    st.success(f"✅ {repo_result['message']}")
                    
                    # Show repo info
                    st.info(format_repo_info_display(repo_result['repo_info']))
                
            elif uploaded_file:
                if uploaded_file.type == "application/pdf":
                    input_text = extract_text_from_pdf(uploaded_file)
                elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                    input_text = extract_text_from_docx(uploaded_file)
                elif uploaded_file.type == "text/csv":
                    uploaded_df = pd.read_csv(uploaded_file)
                    input_text = f"Dataset: {len(uploaded_df)} rows, {len(uploaded_df.columns)} columns\nColumns: {', '.join(uploaded_df.columns)}"
                else:
                    # Handle TXT and PY files
                    input_text = uploaded_file.read().decode('utf-8')
                    
            elif text_input:
                input_text = text_input
            else:
                st.error("Please provide input: upload a file, paste text, or enter a GitHub URL")
                return
            
            if not input_text or len(input_text) < 50:
                st.error("Input text is too short. Please provide more content.")
                return
            
            # Auto-detect artifact type
            artifact_type = "Code" if is_github_repo else detect_artifact_type(input_text, uploaded_df)
            
            # Run analysis
            with st.spinner("🔍 Analyzing your project..."):
                progress = st.progress(0)
                status = st.empty()
                
                # === IP & COPYRIGHT CHECK (only if enabled) ===
                ip_copyright_findings = []
                if check_options.get('copyright', False):
                    status.text("Running IP & Copyright analysis...")
                    progress.progress(8)
                    
                    try:
                        ip_result = analyze_ip_copyright(input_text)
                        if ip_result["status"] == "success":
                            ip_copyright_findings = ip_result["findings"]
                    except Exception as e:
                        st.warning(f"IP & Copyright check encountered an error: {str(e)}")
                
                # === PRIVACY CHECK (only if enabled) ===
                privacy_findings = []
                if check_options.get('privacy', False):
                    status.text("Running Privacy & PII analysis...")
                    progress.progress(18)
                    
                    try:
                        privacy_input_type = "code" if artifact_type == "Code" else "text"
                        privacy_result = run_privacy_check(input_text, privacy_input_type)
                        
                        if privacy_result["status"] == "success":
                            privacy_findings = privacy_result["findings"]
                    except Exception as e:
                        st.warning(f"Privacy check encountered an error: {str(e)}")
                
                # Code security (always run for code)
                status.text("Checking code security...")
                progress.progress(28)
                code_issues = detect_code_issues(input_text)
                
                # Ethical keywords
                status.text("Scanning for ethical keywords...")
                progress.progress(35)
                keyword_flags = check_ethical_keywords(input_text)
                
                # === BIAS CHECK (only if enabled) ===
                bias_findings = []
                if check_options.get('bias', False):
                    status.text("Running bias & fairness analysis...")
                    progress.progress(45)
                    
                    if uploaded_df is not None:
                        bias_result = run_bias_check(uploaded_df, "dataset")
                    elif artifact_type == "Code":
                        bias_result = run_bias_check(input_text, "code")
                    else:
                        bias_result = run_bias_check(input_text, "text")
                    
                    if bias_result["status"] == "success":
                        bias_findings = bias_result["findings"]
                
                # === PLAGIARISM CHECK (only if enabled) ===
                plagiarism_findings = []
                if check_options.get('plagiarism', False):
                    status.text("Running plagiarism check (this may take a moment)...")
                    progress.progress(55)
                    
                    try:
                        def plag_progress(current, total, phrase):
                            pct = 55 + int((current / total) * 25)
                            progress.progress(min(pct, 80))
                            status.text(f"Checking phrase {current}/{total}...")
                        
                        plagiarism_result = run_plagiarism_check(
                            input_text, 
                            max_phrases=8,
                            progress_callback=plag_progress
                        )
                        
                        if plagiarism_result["status"] == "success":
                            plagiarism_findings = plagiarism_result["findings"]
                            st.session_state.plagiarism_percentage = plagiarism_result.get("plagiarism_percentage", 0)
                            st.session_state.plagiarism_sources = plagiarism_result.get("sources", [])
                    except Exception as e:
                        st.warning(f"Plagiarism check encountered an error: {str(e)}")
                
                # Main Groq analysis
                status.text("Running AI analysis...")
                progress.progress(85)
                results = analyze_with_groq(input_text, artifact_type, check_options, uploaded_df)
                
                progress.progress(100)
                status.text("Analysis complete!")
                
                if results:
                    # Add IP & Copyright findings (only if checked)
                    if check_options.get('copyright', False) and ip_copyright_findings:
                        results['issues'].extend(ip_copyright_findings)
                        high_ip = [f for f in ip_copyright_findings if f['severity'] == 'high']
                        if high_ip and results.get('overall_score') != 'high':
                            results['overall_score'] = 'high'
                    
                    # Add Privacy findings (only if checked)
                    if check_options.get('privacy', False) and privacy_findings:
                        results['issues'].extend(privacy_findings)
                        high_privacy = [f for f in privacy_findings if f['severity'] == 'high']
                        if high_privacy and results.get('overall_score') != 'high':
                            results['overall_score'] = 'high'
                    
                    # Add Bias findings (only if checked)
                    if check_options.get('bias', False) and bias_findings:
                        results['issues'].extend(bias_findings)
                        high_bias = [f for f in bias_findings if f['severity'] == 'high']
                        if high_bias and results.get('overall_score') != 'high':
                            results['overall_score'] = 'high'
                    
                    # Add Plagiarism findings (only if checked)
                    if check_options.get('plagiarism', False) and plagiarism_findings:
                        results['issues'].extend(plagiarism_findings)
                        high_plag = [f for f in plagiarism_findings if f['severity'] == 'high']
                        if high_plag and results.get('overall_score') != 'high':
                            results['overall_score'] = 'high'
                    
                    # Sync plagiarism score with AI detection
                    if check_options.get('plagiarism', False):
                        ai_plag_issues = [i for i in results.get('issues', []) 
                                         if i.get('category') == 'Plagiarism' and i.get('severity') == 'high']
                        current_plag_pct = st.session_state.get('plagiarism_percentage', 0)
                        if ai_plag_issues and current_plag_pct < 20:
                            st.session_state.plagiarism_percentage = max(current_plag_pct, 50.0)
                    
                    # Add code security findings
                    if code_issues:
                        for issue in code_issues:
                            results['issues'].append({
                                'category': 'Security',
                                'severity': issue['severity'],
                                'title': issue['type'],
                                'evidence': f"Detected {issue['count']} occurrence(s)",
                                'location': 'Code sections',
                                'recommendation': 'Avoid using dangerous operations that can execute arbitrary code',
                                'suggested_rewrite': 'Use safer alternatives like subprocess.run() with shell=False'
                            })
                    
                    # Save to session state
                    st.session_state.analysis_results = results
                    st.session_state.bias_findings = bias_findings
                    st.session_state.ip_findings = ip_copyright_findings
                    st.session_state.privacy_findings = privacy_findings
                    st.session_state.plagiarism_findings = plagiarism_findings
                    
                    # Clear progress
                    progress.empty()
                    status.empty()
                    
                    # Show balloons
                    st.balloons()
                    
                    # Show success message
                    st.success("✅ Analysis complete! Switch to the **Results** tab to view your report ➡️")
                else:
                    st.error("Analysis failed. Please check your Groq API key and try again.")
                    
    with tab2:
        if st.session_state.analysis_results:
            bias_findings = st.session_state.get('bias_findings', [])
            ip_findings = st.session_state.get('ip_findings', [])
            privacy_findings = st.session_state.get('privacy_findings', [])
            plagiarism_findings = st.session_state.get('plagiarism_findings', [])
            
            render_results(
                st.session_state.analysis_results, 
                bias_findings, 
                ip_findings, 
                privacy_findings,
                plagiarism_findings
            )
            
            # Show plagiarism sources if available
            if st.session_state.get('plagiarism_sources'):
                with st.expander("🌐 Plagiarism Sources Found"):
                    sources = st.session_state.plagiarism_sources
                    for i, source in enumerate(sources[:5], 1):
                        st.markdown(f"**{i}. {source.get('title', 'Unknown Source')[:60]}**")
                        st.caption(f"URL: {source.get('url', 'N/A')[:80]}")
                        if source.get('phrases_matched'):
                            st.caption(f"Matched phrases: {len(source['phrases_matched'])}")
                        st.markdown("---")
        else:
            st.info("No analysis results yet. Upload and analyze a project first.")

if __name__ == "__main__":
    main()